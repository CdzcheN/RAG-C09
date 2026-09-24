#!/usr/bin/env python3
"""把 docs/开题报告.md 转成 docs/开题报告.docx（不依赖 pandoc）。

支持：标题层级（#/##/###）、**粗体** 与 `行内代码`、pipe 表格、有序/无序列表、
引用块、图片（![alt](path) → 居中插图 + 图注）、参考文献悬挂缩进、页脚页码、A4 与中文字体。

用法:
  pip install python-docx          # 可选依赖，仅生成本文档时需要
  python scripts/md2docx.py

说明: 交付前建议用 Word 或 LibreOffice 打开检查排版；
      也可用 libreoffice --headless --convert-to pdf 生成 PDF 预览。
"""
from __future__ import annotations

import pathlib
import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("[fail] 缺少 python-docx，请先执行：python -m pip install python-docx", file=sys.stderr)
    raise SystemExit(2)

ROOT = pathlib.Path(__file__).resolve().parent.parent
MD = ROOT / "docs" / "开题报告.md"
DOCX = ROOT / "docs" / "开题报告.docx"

BODY_ASCII, BODY_CJK = "Times New Roman", "Noto Serif CJK SC"
HEAD_ASCII, HEAD_CJK = "Arial", "Noto Sans CJK SC"
MONO = "DejaVu Sans Mono"
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
IMG = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
REF = re.compile(r"^\[\d+\]\s")


def set_run_fonts(run, ascii_font: str, cjk_font: str) -> None:
    run.font.name = ascii_font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cjk_font)


def style_font(style, ascii_font: str, cjk_font: str, size_pt: float | None = None,
               color: str | None = None) -> None:
    style.font.name = ascii_font
    if size_pt:
        style.font.size = Pt(size_pt)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cjk_font)


def add_rich(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = MONO
            continue
        else:
            run = par.add_run(part)
        set_run_fonts(run, BODY_ASCII, BODY_CJK)


def add_page_number(doc) -> None:
    par = doc.sections[0].footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, "PAGE"),
                             ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)
    set_run_fonts(run, BODY_ASCII, BODY_CJK)


def build() -> dict:
    md_lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for attr, val in (("top_margin", 2.5), ("bottom_margin", 2.5), ("left_margin", 2.6), ("right_margin", 2.6)):
        setattr(sec, attr, Cm(val))

    style_font(doc.styles["Normal"], BODY_ASCII, BODY_CJK, 10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.35
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    style_font(doc.styles["Title"], HEAD_ASCII, HEAD_CJK, 18, "1C2B3A")
    for lvl, size in ((1, 14), (2, 12), (3, 11)):
        style_font(doc.styles[f"Heading {lvl}"], HEAD_ASCII, HEAD_CJK, size, "1C2B3A")

    i, n_lines, stats = 0, len(md_lines), {"head": 0, "tab": 0, "img": 0, "para": 0, "ref": 0}
    while i < n_lines:
        line = md_lines[i].rstrip()

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        m = IMG.match(line)
        if m:
            alt, rel = m.group(1), m.group(2)
            path = (MD.parent / rel).resolve()
            if path.exists():
                pic_par = doc.add_paragraph()
                pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic_par.add_run().add_picture(str(path), width=Cm(15.0))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(alt)
                run.font.size = Pt(9)
                run.italic = True
                set_run_fonts(run, BODY_ASCII, BODY_CJK)
                stats["img"] += 1
            i += 1
            continue

        if line.startswith("|"):
            block = []
            while i < n_lines and md_lines[i].lstrip().startswith("|"):
                block.append(md_lines[i])
                i += 1
            rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in block]
            rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row[: len(rows[0])]):
                    par = table.cell(ri, ci).paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    add_rich(par, cell_text)
                    for run in par.runs:
                        run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True
            stats["tab"] += 1
            doc.add_paragraph()
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line[level:].strip()
            if level == 1:
                doc.add_paragraph(text, style="Title")
                if i + 1 < n_lines and md_lines[i + 1].strip().startswith("**开题报告**"):
                    sub = doc.add_paragraph()
                    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_rich(sub, md_lines[i + 1].strip())
                    for run in sub.runs:
                        run.font.size = Pt(11)
                    i += 1
            else:
                doc.add_heading(text, level=min(level - 1, 3))
            stats["head"] += 1
            i += 1
            continue

        if line.lstrip().startswith(("- ", "* ")):
            indent = (len(line) - len(line.lstrip())) // 2
            par = doc.add_paragraph(style="List Bullet" if indent < 1 else "List Bullet 2")
            add_rich(par, line.lstrip()[2:].strip())
            stats["para"] += 1
            i += 1
            continue

        if re.match(r"^\d+\.\s", line.strip()):
            par = doc.add_paragraph(style="List Number")
            add_rich(par, re.sub(r"^\d+\.\s", "", line.strip()))
            stats["para"] += 1
            i += 1
            continue

        if line.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich(par, line.lstrip("> ").strip())
            for run in par.runs:
                run.font.size = Pt(11)
                run.bold = True
            stats["para"] += 1
            i += 1
            continue

        par = doc.add_paragraph()
        add_rich(par, line.strip())
        if REF.match(line.strip()):
            par.paragraph_format.left_indent = Cm(0.85)
            par.paragraph_format.first_line_indent = Cm(-0.85)
            par.paragraph_format.space_after = Pt(3)
            stats["ref"] += 1
        else:
            stats["para"] += 1
        i += 1

    add_page_number(doc)
    doc.save(DOCX)
    return stats


def main() -> int:
    if not MD.exists():
        print(f"[fail] 未找到源文件 {MD}", file=sys.stderr)
        return 1
    stats = build()
    print(f"[ok] {DOCX.relative_to(ROOT)}  ({DOCX.stat().st_size} bytes)")
    print(f"     标题 {stats['head']} · 表格 {stats['tab']} · 插图 {stats['img']} · "
          f"段落 {stats['para']} · 参考文献 {stats['ref']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
